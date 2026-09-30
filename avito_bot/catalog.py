"""Чтение таблицы-каталога (xlsx или опубликованная Google Таблица) и дерево кнопок."""
from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass, field
from urllib.request import urlopen

DEFAULTS = {
    "город по умолчанию": "Санкт-Петербург + ЛО (50 км)",
    "радиус поиска, км": "50",
    "состояние по умолчанию": "б/у",
    "только с доставкой по умолчанию": "нет",
    "стоп-слова для всех строк (только исправное)": "",
    "минус-слова в заголовке для всех строк": ("коробка, коробку, коробки, короб, упаковка, бокс, box, кронштейн, держатель, "
                                              "подставка, кабель, переходник, райзер, riser, бэкплейт, backplate, "
                                              "кулер для, вентилятор для, вентиляторы, радиатор для, термопрокладки, "
                                              "водоблок, куплю, обмен, скупка, выкуп"),
    "сколько объявлений показывать в боте": "10",
    "сортировка": "по цене",
    "интервал парсинга, минут": "30",
    "уведомлять о новых объявлениях": "да",
}


def split_words(value) -> list[str]:
    if not value:
        return []
    return [w.strip() for w in str(value).split(",") if w.strip()]


def to_int(value) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):  # из xlsx числа приходят как 16000.0
        return int(value)
    text = re.sub(r"[\s\u00a0\u202f₽р.]*$", "", str(value).strip().lower())
    text = re.sub(r"[.,]\d{1,2}$", "", text)  # копейки / «16000.0»
    digits = "".join(ch for ch in text if ch.isdigit())
    return int(digits) if digits else None


@dataclass
class Item:
    path: tuple[str, ...]          # разделы + модель, т.е. путь кнопок
    query: str
    price_min: int | None = None
    price_max: int | None = None
    must: list[str] = field(default_factory=list)
    minus: list[str] = field(default_factory=list)
    city: str = ""
    condition: str = ""
    delivery_only: bool = False
    variant: str = ""               # версия модели (8GB / 16GB): своя строка и цена, но общая кнопка

    @property
    def model(self) -> str:
        return self.path[-1]

    @property
    def title(self) -> str:
        return f"{self.model} {self.variant}".strip()

    @property
    def key(self) -> str:
        return node_id(self.path + ((self.variant,) if self.variant else ()))

    @property
    def category(self) -> str:
        return self.path[0]


@dataclass
class Settings:
    raw: dict[str, str]

    def get(self, name: str) -> str:
        return str(self.raw.get(name, DEFAULTS.get(name, "")) or DEFAULTS.get(name, ""))

    @property
    def city(self) -> str: return self.get("город по умолчанию")
    @property
    def spb_only(self) -> bool:
        return "россия" not in self.city.lower()

    @property
    def radius_km(self) -> int: return to_int(self.get("радиус поиска, км")) or 0
    @property
    def condition(self) -> str: return self.get("состояние по умолчанию")
    @property
    def delivery_only(self) -> bool: return self.get("только с доставкой по умолчанию").lower() == "да"
    @property
    def stop_words(self) -> list[str]: return split_words(self.get("стоп-слова для всех строк (только исправное)"))
    @property
    def title_stop_words(self) -> list[str]: return split_words(self.get("минус-слова в заголовке для всех строк"))
    @property
    def show_count(self) -> int: return to_int(self.get("сколько объявлений показывать в боте")) or 10
    @property
    def sort_by_price(self) -> bool: return "цен" in self.get("сортировка").lower()
    @property
    def interval_min(self) -> int: return max(5, to_int(self.get("интервал парсинга, минут")) or 30)
    @property
    def notify(self) -> bool: return self.get("уведомлять о новых объявлениях").lower() == "да"

    def avito_section(self, category: str) -> str | None:
        return self.raw.get(f"раздел авито: {category.lower()}") or None


def node_id(path) -> str:
    """Короткий стабильный id узла для callback_data (лимит Telegram 64 байта)."""
    return hashlib.sha1("\x1f".join(path).encode()).hexdigest()[:12]


def _norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # Google отдаёт число 30 как 30.0, а в соседней строке «30» текстом: без этого два раздела
    return str(v).strip()


# Колонки листа «Каталог» (по заголовкам, чтобы порядок можно было менять)
COL = {
    "path": ["раздел 1", "раздел 2", "раздел 3", "раздел 4", "раздел 5"],
    "model": "модель (кнопка)", "variant": "вариант", "on": "вкл",
    "pmin": "цена от, ₽", "pmax": "цена до, ₽",
    "query": "поисковый запрос", "must": "обязательные слова", "minus": "минус-слова (доп.)",
    "city": "город", "cond": "состояние", "delivery": "только с доставкой",
}


def parse_rows(rows: list[list]) -> list[Item]:
    if not rows:
        return []
    header = [_norm(h).lower() for h in rows[0]]
    idx = {h: i for i, h in enumerate(header)}

    def cell(r, name):
        i = idx.get(name)
        return _norm(r[i]) if i is not None and i < len(r) else ""

    items = []
    for r in rows[1:]:
        model = cell(r, COL["model"])
        if not model or cell(r, COL["on"]).lower() == "нет":
            continue
        path = tuple(p for p in (cell(r, c) for c in COL["path"]) if p) + (model,)
        items.append(Item(
            path=path,
            query=cell(r, COL["query"]) or model,
            price_min=to_int(cell(r, COL["pmin"])),
            price_max=to_int(cell(r, COL["pmax"])),
            must=split_words(cell(r, COL["must"])),
            minus=split_words(cell(r, COL["minus"])),
            city=cell(r, COL["city"]),
            condition=cell(r, COL["cond"]),
            delivery_only=cell(r, COL["delivery"]).lower() == "да",
            variant=cell(r, COL["variant"]),
        ))
    # Несколько строк с одной моделью = версии одной кнопки. Если «Вариант» не заполнен,
    # называем версию по обязательным словам (например «16GB»)
    groups: dict[tuple, list[Item]] = {}
    for it in items:
        groups.setdefault(it.path, []).append(it)
    for group in groups.values():
        if len(group) > 1:
            seen = set()
            for i, it in enumerate(group, 1):
                if not it.variant:
                    it.variant = ", ".join(it.must) or f"#{i}"
                if it.variant in seen:
                    it.variant += f" #{i}"
                seen.add(it.variant)
    return items


def parse_settings(rows: list[list]) -> Settings:
    raw = {}
    for r in rows[1:]:
        if len(r) >= 2 and _norm(r[0]):
            raw[_norm(r[0]).lower()] = _norm(r[1])
    return Settings(raw)


def _read_xlsx(path) -> tuple[list[list], list[list]]:
    """Листы ищем по заголовкам, а не по названию: при импорте в Google названия могут поменяться."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    cat, st = None, []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        head = [_norm(h).lower() for h in (rows[0] if rows else [])]
        if cat is None and COL["model"] in head:
            cat = rows
        elif head[:1] == ["параметр"]:
            st = rows
    if cat is None:
        names = ", ".join(wb.sheetnames)
        raise RuntimeError(f"Не нашёл лист с колонкой «Модель (кнопка)». Листы в таблице: {names}")
    return cat, st


def _google_xlsx_url(url: str) -> str:
    # https://docs.google.com/spreadsheets/d/<id>/edit... -> выгрузка всей таблицы в xlsx (все листы разом)
    doc_id = url.split("/d/")[1].split("/")[0]
    return f"https://docs.google.com/spreadsheets/d/{doc_id}/export?format=xlsx"


def load(source: str) -> tuple[list[Item], Settings]:
    """source: путь к .xlsx или ссылка на Google Таблицу (доступ «все, у кого есть ссылка»)."""
    if source.startswith("http"):
        with urlopen(_google_xlsx_url(source), timeout=30) as resp:
            data = resp.read()
        if not data.startswith(b"PK"):
            raise RuntimeError("Google не отдал таблицу: включи доступ «Все, у кого есть ссылка → Читатель»")
        source = io.BytesIO(data)
    cat, st = _read_xlsx(source)
    return parse_rows(cat), parse_settings(st)


@dataclass
class Node:
    id: str
    title: str
    path: tuple[str, ...]
    children: dict[str, "Node"] = field(default_factory=dict)
    items: list[Item] = field(default_factory=list)
    parent: "Node | None" = None


class Tree:
    def __init__(self, items: list[Item]):
        self.root = Node(id="root", title="Каталог", path=())
        self.by_id: dict[str, Node] = {"root": self.root}
        self.items = {it.key: it for it in items}
        for it in items:
            node = self.root
            for depth in range(1, len(it.path) + 1):
                title = it.path[depth - 1]
                child = node.children.get(title)
                if child is None:
                    sub = it.path[:depth]
                    child = Node(id=node_id(sub), title=title, path=sub, parent=node)
                    node.children[title] = child
                    self.by_id[child.id] = child
                node = child
            node.items.append(it)

    def get(self, nid: str) -> Node | None:
        return self.by_id.get(nid)
